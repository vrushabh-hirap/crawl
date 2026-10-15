"""Resume text extractor.

Delegates to Stage 1 Ingestion pipeline for backward compatibility.
"""

from crawl.jobs.ats.pipeline.stage1_ingestion import ingest


def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract plain text from a résumé file using Stage 1 Ingestion pipeline."""
    result = ingest(file_bytes, filename)
    return result.text



# ── Helpers ──────────────────────────────────────────────────────────────────

def _get_extension(filename: str) -> str:
    """Return the lowercase file extension including the leading dot."""
    idx = filename.rfind(".")
    if idx == -1:
        return ""
    return filename[idx:].lower()


def _extract_pdf(data: bytes) -> str:
    """Extract text from a PDF using pypdf."""
    try:
        import pypdf  # type: ignore
    except ImportError:
        raise ImportError(
            "pypdf is required for PDF parsing. "
            "Install via: pip install crawl-jobs"
        )

    reader = pypdf.PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)

    result = "\n".join(pages).strip()
    if not result:
        raise ValueError(
            "Could not extract any text from the PDF. "
            "The file may be image-only (scanned). Please use a text-based PDF."
        )
    return result


def _extract_docx(data: bytes) -> str:
    """Extract text from a DOCX file using python-docx."""
    try:
        import docx  # type: ignore
    except ImportError:
        raise ImportError(
            "python-docx is required for DOCX parsing. "
            "Install via: pip install crawl-jobs"
        )

    doc = docx.Document(io.BytesIO(data))
    paragraphs = [para.text for para in doc.paragraphs]
    # Also pull text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.append(cell.text)

    result = "\n".join(p for p in paragraphs if p.strip()).strip()
    if not result:
        raise ValueError("Could not extract any text from the DOCX file.")
    return result


def _extract_plain(data: bytes) -> str:
    """Decode plain-text bytes, trying UTF-8 then latin-1 as fallback."""
    try:
        return data.decode("utf-8").strip()
    except UnicodeDecodeError:
        return data.decode("latin-1").strip()
